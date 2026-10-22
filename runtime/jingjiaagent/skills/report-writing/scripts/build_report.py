"""Build editable Word from reviewed report content, without network access."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(spec, output):
    if not isinstance(spec.get("title"), str) or not spec["title"].strip(): raise ValueError("title is required")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = document.styles[style_name]; style.font.name = "Noto Sans CJK SC"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(spec["title"], 0)
    if spec.get("subtitle"): document.add_paragraph(spec["subtitle"], style="Subtitle")
    if spec.get("summary"): document.add_paragraph(spec["summary"])
    for section_spec in spec.get("sections", []):
        document.add_heading(section_spec["heading"], 1)
        for paragraph in section_spec.get("paragraphs", []): document.add_paragraph(str(paragraph))
        table_spec = section_spec.get("table")
        if table_spec:
            headers = table_spec["headers"]
            if not headers or len(headers) > 12: raise ValueError("table must have 1–12 columns")
            table = document.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
            for cell, title in zip(table.rows[0].cells, headers): cell.text = str(title)
            repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in table_spec.get("rows", []):
                if len(row) != len(headers): raise ValueError("table row width differs from headers")
                for cell, value in zip(table.add_row().cells, row): cell.text = str(value)
    if spec.get("sources"):
        document.add_heading("资料来源", 1)
        for index, source in enumerate(spec["sources"], 1): document.add_paragraph(str(index) + ". " + str(source))
    output.parent.mkdir(parents=True, exist_ok=True); document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path, default=Path(".jingjiaagent/output/report.docx")); args = parser.parse_args()
    build(json.loads(args.input.read_text(encoding="utf-8")), args.output)
