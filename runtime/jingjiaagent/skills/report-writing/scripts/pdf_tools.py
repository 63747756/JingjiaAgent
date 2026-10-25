"""Offline PDF text/table extraction and page assembly; never overwrite files."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import sys

import pdfplumber
from pypdf import PdfReader, PdfWriter


MAX_BYTES = 32 << 20
MAX_PAGES = 200
MAX_CELLS = 200_000


class PdfToolError(ValueError):
    """An actionable input or output error."""


class LimitedBuffer(io.BytesIO):
    def write(self, data):
        if self.tell() + len(data) > MAX_BYTES:
            raise PdfToolError("输出超过 32 MiB，请减少页数后重试")
        return super().write(data)


def read_pdf(source):
    source = Path(source)
    try:
        with source.open("rb") as stream:
            data = stream.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            raise PdfToolError("输入文件超过 32 MiB")
        if not data.startswith(b"%PDF-"):
            raise PdfToolError("不是有效的 PDF 文件")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise PdfToolError("PDF 已加密，请先提供解密后的副本")
        if not 1 <= len(reader.pages) <= MAX_PAGES:
            raise PdfToolError("PDF 页数必须为 1–200 页")
        return source, data, reader
    except PdfToolError:
        raise
    except OSError as error:
        raise PdfToolError(f"无法读取 PDF {source.name}: {error}") from error
    except Exception as error:
        raise PdfToolError(f"PDF 无法解析或已经损坏: {source.name} ({error})") from error


def page_numbers(selection, total):
    if selection is None:
        return list(range(1, total + 1))
    result = []
    for part in selection.split(","):
        part = part.strip()
        if not re.fullmatch(r"[1-9][0-9]*(?:-[1-9][0-9]*)?", part):
            raise PdfToolError("页码格式错误，请使用 1,3-5 等从 1 开始的页码")
        edges = [int(value) for value in part.split("-")]
        start, end = edges[0], edges[-1]
        if start > end or end > total:
            raise PdfToolError(f"页码范围无效：{part}，文件共 {total} 页")
        for page in range(start, end + 1):
            if page in result:
                raise PdfToolError(f"页码重复：{page}")
            result.append(page)
    return result


def save_new(output, data):
    output = Path(output)
    if len(data) > MAX_BYTES:
        raise PdfToolError("输出超过 32 MiB，请减少页数后重试")
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        with output.open("xb") as stream:
            try:
                stream.write(data)
            except Exception:
                stream.close()
                output.unlink(missing_ok=True)
                raise
    except FileExistsError as error:
        raise PdfToolError(f"输出文件已存在，禁止覆盖原件或已有成果：{output}") from error
    except OSError as error:
        raise PdfToolError(f"无法保存输出：{error}") from error


def extract_text(source, output, selection=None, output_format="json"):
    source, data, reader = read_pdf(source)
    result = {"source_file": source.name, "source_sha256": hashlib.sha256(data).hexdigest(),
              "total_pages": len(reader.pages), "pages": [], "warnings": []}
    extracted_bytes = 0
    for number in page_numbers(selection, len(reader.pages)):
        try:
            text = reader.pages[number - 1].extract_text() or ""
        except Exception as error:
            raise PdfToolError(f"第 {number} 页文字提取失败：{error}") from error
        extracted_bytes += len(text.encode("utf-8"))
        if extracted_bytes > MAX_BYTES:
            raise PdfToolError("提取文字超过 32 MiB，请减少页数后重试")
        result["pages"].append({"page": number, "text": text})
        if not text.strip():
            result["warnings"].append(f"第 {number} 页没有可提取文字，可能为扫描件或空白页；扫描件需要 OCR。")
    if output_format == "json":
        encoded = json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")
    elif output_format == "text":
        chunks = [f"来源：{source.name}\nSHA256：{result['source_sha256']}\n"]
        chunks.extend(f"\n--- 第 {page['page']} 页 ---\n{page['text']}\n" for page in result["pages"])
        chunks.extend(f"\n警告：{warning}\n" for warning in result["warnings"])
        encoded = "".join(chunks).encode("utf-8")
    else:
        raise PdfToolError("文字输出格式必须为 json 或 text")
    save_new(output, encoded)
    return result


def safe_cell(value):
    # PDF text must not become a spreadsheet formula when opened by an employee.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def extract_tables(source, output, selection=None):
    source, data, reader = read_pdf(source)
    output = Path(output)
    if output.suffix.lower() not in (".csv", ".xlsx"):
        raise PdfToolError("表格输出文件必须为 .csv 或 .xlsx")
    selected = page_numbers(selection, len(reader.pages))
    rows, warnings = [], []
    max_columns = cell_count = extracted_bytes = table_count = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as document:
            for number in selected:
                page = document.pages[number - 1]
                try:
                    tables = page.extract_tables()
                    if not tables:
                        warnings.append(f"第 {number} 页未检测到规则表格；无边框、合并单元格或扫描表格需要人工复核。")
                    for table_index, table in enumerate(tables, 1):
                        table_count += 1
                        for row_index, values in enumerate(table, 1):
                            cell_count += len(values)
                            extracted_bytes += sum(len((value or "").encode("utf-8")) for value in values)
                            if cell_count > MAX_CELLS or extracted_bytes > MAX_BYTES:
                                raise PdfToolError("表格内容超过限制，请减少页数后重试")
                            if any(value is not None and len(value) > 32767 for value in values):
                                raise PdfToolError("单元格文字超过 Excel 的 32767 字符限制，请改用文字提取")
                            max_columns = max(max_columns, len(values))
                            rows.append([source.name, number, table_index, row_index, *values])
                finally:
                    page.close()
    except PdfToolError:
        raise
    except Exception as error:
        raise PdfToolError(f"PDF 表格提取失败：{error}") from error
    headers = ["source_file", "source_page", "source_table", "source_row"]
    headers.extend(f"column_{index}" for index in range(1, max_columns + 1))
    buffer = LimitedBuffer()
    if output.suffix.lower() == ".csv":
        text_stream = io.TextIOWrapper(buffer, encoding="utf-8-sig", newline="")
        writer = csv.writer(text_stream)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([safe_cell(value) for value in row] + [None] * (len(headers) - len(row)))
        text_stream.flush()
        text_stream.detach()
    else:
        from openpyxl import Workbook

        book = Workbook(write_only=True)
        sheet = book.create_sheet("tables")
        sheet.freeze_panes = "E2"
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(value) for value in row])
        metadata = book.create_sheet("source")
        metadata.append(["source_file", safe_cell(source.name)])
        metadata.append(["source_sha256", hashlib.sha256(data).hexdigest()])
        metadata.append(["pages", ",".join(str(page) for page in selected)])
        for warning in warnings:
            metadata.append(["warning", warning])
        book.save(buffer)
    save_new(output, buffer.getvalue())
    return {"source_file": source.name, "tables": table_count, "rows": len(rows), "warnings": warnings}


def reject_forms_for_page_assembly(reader):
    # add_page copies page objects, but does not preserve the document-level form.
    form = reader.trailer["/Root"].get("/AcroForm")
    if form is None:
        return
    form = form.get_object()
    fields = form.get("/Fields")
    if fields is not None:
        fields = fields.get_object()
    if fields or "/XFA" in form:
        raise PdfToolError(
            "PDF 包含交互式表单（AcroForm/XFA），暂不支持页面提取或合并。"
            "请先另存或打印为普通 PDF 副本；仍可使用文字和表格提取。"
        )


def extract_pages(source, output, selection):
    _, _, reader = read_pdf(source)
    reject_forms_for_page_assembly(reader)
    selected = page_numbers(selection, len(reader.pages))
    writer = PdfWriter()
    for number in selected:
        writer.add_page(reader.pages[number - 1])
    buffer = LimitedBuffer()
    writer.write(buffer)
    save_new(output, buffer.getvalue())
    return {"pages": selected, "warnings": []}


def merge_pdfs(sources, output):
    if not sources:
        raise PdfToolError("至少需要一个输入 PDF")
    writer = PdfWriter()
    total_bytes = total_pages = 0
    for source in sources:
        _, data, reader = read_pdf(source)
        reject_forms_for_page_assembly(reader)
        total_bytes += len(data)
        total_pages += len(reader.pages)
        if total_bytes > MAX_BYTES or total_pages > MAX_PAGES:
            raise PdfToolError("合并输入总量不得超过 32 MiB 或 200 页")
        for page in reader.pages:
            writer.add_page(page)
    buffer = LimitedBuffer()
    writer.write(buffer)
    save_new(output, buffer.getvalue())
    return {"pages": total_pages, "warnings": []}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("text", "tables", "extract", "merge"):
        command = commands.add_parser(name)
        command.add_argument("input", type=Path, nargs="+" if name == "merge" else None)
        command.add_argument("--output", type=Path, required=True)
        if name != "merge":
            command.add_argument("--pages", required=name == "extract", help="1-based page selection, e.g. 1,3-5")
        if name == "text":
            command.add_argument("--format", choices=("json", "text"), default="json")
    args = parser.parse_args(argv)
    try:
        if args.command == "text":
            result = extract_text(args.input, args.output, args.pages, args.format)
        elif args.command == "tables":
            result = extract_tables(args.input, args.output, args.pages)
        elif args.command == "extract":
            result = extract_pages(args.input, args.output, args.pages)
        else:
            result = merge_pdfs(args.input, args.output)
        for warning in result.get("warnings", []):
            print(f"警告：{warning}", file=sys.stderr)
        print(json.dumps({"output": str(args.output), **{key: value for key, value in result.items() if key != "pages"}},
                         ensure_ascii=False))
        return 0
    except Exception as error:
        print(f"PDF 处理失败：{error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
