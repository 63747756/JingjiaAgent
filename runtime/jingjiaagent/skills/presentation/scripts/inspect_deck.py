"""Inspect an existing PPTX without modifying it; indices address python-pptx shapes."""
import argparse
import json
from pathlib import Path
import sys
import zipfile


MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 4096
MAX_SLIDES = 500
MAX_SHAPES = 10000
MAX_TABLE_CELLS = 20000
MAX_TEXT_CHARS = 2 * 1024 * 1024
MAX_GROUP_DEPTH = 20


class InspectionError(ValueError):
    """An unsupported, malformed, or excessively large presentation."""


def validate_package(source):
    if source.suffix.lower() != ".pptx":
        raise InspectionError("仅支持 .pptx；请先将旧版 .ppt 转换为 .pptx。")
    if not source.is_file():
        raise InspectionError("输入文件不存在或不是普通文件。")
    if source.stat().st_size > MAX_FILE_BYTES:
        raise InspectionError("PPTX 超过 64 MiB 大小限制，请先拆分。")
    try:
        with zipfile.ZipFile(source) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise InspectionError("PPTX 包含过多文件，请先拆分。")
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise InspectionError("PPTX 解压后超过 256 MiB 限制，请先拆分。")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise InspectionError("PPTX 包中存在重复文件。")
            if any(entry.flag_bits & 1 for entry in entries):
                raise InspectionError("不支持加密的 PPTX 包，请先解密。")
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(names):
                raise InspectionError("文件不是有效的 PPTX 演示文稿。")
    except zipfile.BadZipFile as exc:
        raise InspectionError("PPTX 包损坏或已加密，无法读取。") from exc


class Budget:
    def __init__(self):
        self.shapes = 0
        self.cells = 0
        self.characters = 0

    def text(self, value):
        value = str(value)
        self.characters += len(value)
        if self.characters > MAX_TEXT_CHARS:
            raise InspectionError("文本内容过多，请使用 --slide 仅检查指定页。")
        return value


def describe_shapes(shapes, budget, parent_path=()):
    if len(parent_path) > MAX_GROUP_DEPTH:
        raise InspectionError("形状分组嵌套过深，无法安全检查。")
    result = []
    for index, shape in enumerate(shapes):
        budget.shapes += 1
        if budget.shapes > MAX_SHAPES:
            raise InspectionError("形状数量过多，请使用 --slide 仅检查指定页。")
        path = (*parent_path, index)
        shape_type = shape.shape_type
        item = {
            "shape_path": list(path),
            "shape_id": shape.shape_id,
            "name": budget.text(shape.name),
            "type": shape_type.name if shape_type is not None else "UNSUPPORTED",
            "bounds_emu": {
                "left": shape.left, "top": shape.top,
                "width": shape.width, "height": shape.height,
            },
        }
        if shape_type is None:
            # python-pptx returns None for valid GraphicFrames such as SmartArt.
            item["warnings"] = ["该对象类型暂不支持解析（如 SmartArt）；仅提供定位和尺寸，内容未完整解析。"]
        if shape.has_text_frame:
            item["text"] = budget.text(shape.text)
        if shape.has_table:
            table = shape.table
            budget.cells += len(table.rows) * len(table.columns)
            if budget.cells > MAX_TABLE_CELLS:
                raise InspectionError("表格单元格过多，请使用 --slide 仅检查指定页。")
            item["table"] = {
                "rows": len(table.rows), "columns": len(table.columns),
                "cells": [
                    [
                        {"text": budget.text(cell.text),
                         "is_merge_origin": cell.is_merge_origin,
                         "is_spanned": cell.is_spanned}
                        for cell in row.cells
                    ] for row in table.rows
                ],
            }
        if shape.has_chart:
            chart = shape.chart
            try:
                item["chart"] = {
                    "type": chart.chart_type.name,
                    "series_names": [budget.text(series.name) for series in chart.series],
                }
            except ValueError as exc:
                # The pinned library rejects some valid OOXML plots (e.g. pie3DChart).
                # Do not downgrade malformed chart values or our own budget errors.
                if not str(exc).startswith("unsupported plot type "):
                    raise
                item["chart"] = {"type": "UNSUPPORTED"}
                item.setdefault("warnings", []).append(
                    "该图表类型暂不支持解析；仅提供定位、尺寸及可读取的标题，数据未完整解析。")
            if chart.has_title and chart.chart_title.has_text_frame:
                item["chart"]["title"] = budget.text(chart.chart_title.text_frame.text)
        if hasattr(shape, "shapes"):
            item["children"] = describe_shapes(shape.shapes, budget, path)
        result.append(item)
    return result


def inspect_deck(source, slide_number=None):
    """Return JSON-compatible metadata. Slide numbers are 1-based; shape paths 0-based."""
    source = Path(source)
    validate_package(source)
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise InspectionError("缺少 python-pptx，请使用已安装该依赖的办公 Python 环境。") from exc
    try:
        presentation = Presentation(source)
        slide_count = len(presentation.slides)
        if slide_count > MAX_SLIDES:
            raise InspectionError("演示文稿超过 500 页，请先拆分。")
        if slide_number is not None and not 1 <= slide_number <= slide_count:
            raise InspectionError(f"页码超出范围：文稿共 {slide_count} 页，页码从 1 开始。")
        slides = ([(slide_number, presentation.slides[slide_number - 1])]
                  if slide_number is not None else enumerate(presentation.slides, 1))
        budget = Budget()
        return {
            "source": str(source),
            "slide_count": slide_count,
            "slide_size_emu": {"width": presentation.slide_width,
                               "height": presentation.slide_height},
            "indexing": "slide_number 从 1 开始；shape_path 各级索引从 0 开始，嵌套分组通过 .shapes 访问。",
            "slides": [{"slide_number": number,
                        "shapes": describe_shapes(slide.shapes, budget)}
                       for number, slide in slides],
        }
    except InspectionError:
        raise
    except Exception as exc:
        raise InspectionError("PPTX 内容损坏或包含无法解析的结构，请用办公软件重新另存为 PPTX 后重试。") from exc


def main(argv=None):
    parser = argparse.ArgumentParser(description="读取 PPTX 页码、形状索引、文字、表格和图表概览，不修改原件。")
    parser.add_argument("input", type=Path)
    parser.add_argument("--slide", type=int, help="仅输出指定页，页码从 1 开始")
    parser.add_argument("--output", type=Path, help="JSON 文件路径；默认输出至标准输出")
    args = parser.parse_args(argv)
    try:
        if args.output and args.output.resolve() == args.input.resolve():
            raise InspectionError("输出路径不能覆盖输入 PPTX。")
        content = json.dumps(inspect_deck(args.input, args.slide), ensure_ascii=False, indent=2)
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(content + "\n", encoding="utf-8")
        else:
            print(content)
        return 0
    except (InspectionError, OSError) as exc:
        print(f"PPT 检查失败：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
