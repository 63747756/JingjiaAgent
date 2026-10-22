"""Deterministic intranet chat/SSE fixture for office product acceptance only."""
import json
import os
import subprocess
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
import time
from pathlib import Path
from urllib.parse import unquote, urlsplit

CODE_SCRIPT = """set -e
mkdir -p .jingjiaagent/work/code .jingjiaagent/output
/opt/office/venv/bin/python - <<'PY'
from pathlib import Path
import subprocess,hashlib,json
root=Path('.')
code='def add(a, b):\\n    return a + b\\n\\nassert add(12, 8) == 20\\n'
target=root/'calculator.py'
target.write_text(code)
subprocess.run(['/opt/office/venv/bin/python',str(target)],check=True)
Path('.jingjiaagent/work/code/calculator.py').write_text(code)
patch=''
if (root/'.git').exists():
 readme=root/'README.md'
 readme.write_text((readme.read_text() if readme.exists() else '')+'\\nUnified assistant coding acceptance.\\n')
 patch=subprocess.run(['git','diff','--binary'],capture_output=True,text=True,check=True).stdout
added=subprocess.run(['git','diff','--no-index','--binary','/dev/null',str(target)],capture_output=True,text=True)
if added.returncode not in (0,1):raise RuntimeError('code patch creation failed')
Path('.jingjiaagent/output/code.patch').write_text(patch+added.stdout)
Path('.jingjiaagent/output/coding-report.md').write_text('# Unified coding acceptance\\n\\nDeterministic simulated model. Actual Python execution verified add(12,8)=20.\\nRepository modification: '+str((root/'.git').exists())+'\\nEditable code and patch are saved.\\n')
Path('.jingjiaagent/work/code/verification.json').write_text(json.dumps({'result':20,'repository':(root/'.git').exists(),'sha256':hashlib.sha256(target.read_bytes()).hexdigest()}))
print('coding-fixture-deliverables-created')
PY
"""

SCRIPT="""set -e
mkdir -p .jingjiaagent/work .jingjiaagent/output
/opt/office/venv/bin/python - <<'PY'
from pathlib import Path
import csv,hashlib,json
from decimal import Decimal
from docx import Document
from openpyxl import Workbook
root=Path('.')
sources=sorted(Path('.jingjiaagent/attachments').rglob('*acceptance.csv'),key=lambda path:(path.stat().st_mtime_ns,str(path))) if Path('.jingjiaagent/attachments').exists() else []
source=sources[-1] if sources else None
rows=[]
if source:
 with source.open(encoding='utf-8-sig',newline='') as stream:
  for row in csv.DictReader(stream):rows.append((row['项目'],Decimal(row['金额'])))
 if not rows:raise ValueError('fixture CSV has no data rows')
else:rows=[('采购',Decimal(12)),('维护',Decimal(8))]
total=sum((value for _,value in rows),Decimal(0))
input_hash=hashlib.sha256(source.read_bytes()).hexdigest() if source else None
origin=source.name if source else '无CSV任务的固定模拟样例'
summary='、'.join(label+format(value,'f') for label,value in rows)+'，合计'+format(total,'f')
doc=Document();doc.add_heading('办公任务验收报告',0);doc.add_paragraph('本报告来自隔离模拟模型，'+summary+'。');doc.add_paragraph('来源：'+origin+'；数据SHA256：'+(input_hash or '固定模拟样例'));doc.save('.jingjiaagent/output/report.docx')
book=Workbook();sheet=book.active;sheet.append(['项目','金额'])
for label,value in rows:sheet.append([label,float(value)])
sheet.append(['合计','=SUM(B2:B'+str(len(rows)+1)+')']);book.save('.jingjiaagent/output/analysis.xlsx')
Path('.jingjiaagent/work/context.json').write_text(json.dumps({'goal':'隔离办公任务验收','completed':['报告','分析表','演示稿'],'source':origin,'input_sha256':input_hash,'source_rows':len(rows),'total':format(total,'f')},ensure_ascii=False))
deck={'title':'模拟资料办公任务验收','slides':[
 {'title':'授权资料与来源','bullets':['隔离模拟模型；使用当前固定的原始资料。','来源：'+origin],'source':origin},
 {'title':'金额分析','bullets':[summary],'chart':{'labels':[label for label,_ in rows],'values':[float(value) for _,value in rows],'series':'金额'},'source':origin},
 {'title':'交付与复核','bullets':['保留可编辑报告、分析表和演示稿。','平台执行公式重算与PDF预览；数据为验收模拟资料。'],'source':origin}]}
Path('.jingjiaagent/work/presentation-data.json').write_text(json.dumps(deck,ensure_ascii=False))
print('office-fixture-deliverables-created')
PY
node /opt/jingjiaagent/skills/presentation/scripts/build_deck.cjs .jingjiaagent/work/presentation-data.json .jingjiaagent/output/presentation.pptx
"""


class Handler(BaseHTTPRequestHandler):
    protocol_version='HTTP/1.1'
    def log_message(self,*_args):pass
    def json(self,value,status=200):
        data=json.dumps(value,ensure_ascii=False).encode();self.send_response(status);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
    def repository(self,body=b''):
        # Read-only smart HTTP is needed because the runtime fetches shallow
        # repository workspaces. This fixture accepts only upload-pack traffic.
        parsed=urlsplit(self.path)
        path=unquote(parsed.path)
        if path not in ('/fixtures/repository.git/info/refs','/fixtures/repository.git/git-upload-pack'):
            return self.json({'error':'not_found'},404)
        if path.endswith('/info/refs') and parsed.query!='service=git-upload-pack':
            return self.json({'error':'read_only_repository'},403)
        env=dict(os.environ,GIT_PROJECT_ROOT='/fixtures',GIT_HTTP_EXPORT_ALL='1',
            PATH_INFO=path.removeprefix('/fixtures'),QUERY_STRING=parsed.query,
            REQUEST_METHOD=self.command,CONTENT_TYPE=self.headers.get('Content-Type',''),
            CONTENT_LENGTH=str(len(body)),REMOTE_ADDR='127.0.0.1',
            GIT_CONFIG_COUNT='2',GIT_CONFIG_KEY_0='safe.directory',GIT_CONFIG_VALUE_0='/fixtures/repository.git',
            GIT_CONFIG_KEY_1='http.receivepack',GIT_CONFIG_VALUE_1='false')
        if self.headers.get('Git-Protocol'):env['HTTP_GIT_PROTOCOL']=self.headers['Git-Protocol']
        process=subprocess.run(['git','http-backend'],input=body,capture_output=True,env=env,timeout=30)
        headers,sep,data=process.stdout.partition(b'\r\n\r\n')
        if not sep:return self.json({'error':'repository_backend_failed'},502)
        fields=[line.decode().partition(':') for line in headers.split(b'\r\n')]
        status=next((int(value.strip().split()[0]) for key,_,value in fields if key.lower()=='status'),200)
        self.send_response(status)
        for key,_,value in fields:
            if key.lower() not in ('status','content-length'):self.send_header(key,value.strip())
        self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
    def do_GET(self):
        if self.path in ('/health','/v1/models','/models'):
            return self.json({'object':'list','data':[{'id':'office-fixture','object':'model','owned_by':'isolated-fixture'}]})
        path=unquote(urlsplit(self.path).path)
        if path.startswith('/fixtures/repository.git/'):
            return self.repository()
        return self.json({'error':'not_found'},404)
    def do_POST(self):
        size=int(self.headers.get('Content-Length','0'))
        if not 0<size<=16<<20:return self.json({'error':'input_limit'},413)
        body=self.rfile.read(size)
        if urlsplit(self.path).path.startswith('/fixtures/repository.git/'):
            return self.repository(body)
        payload=json.loads(body)
        # Reserved acceptance models keep a round open while the UI saves the
        # next selection. No ordinary fixture or configured real model is delayed.
        if '-next-turn-hold-' in payload.get('model',''):
            time.sleep(4)
        request_path=urlsplit(self.path).path
        if request_path in ('/v1/responses','/responses'):
            text='隔离协议验收响应：已收到本轮请求。此响应来自模拟模型。'
            item={'id':'msg_fixture_'+str(time.time_ns()),'type':'message','role':'assistant','status':'completed','content':[{'type':'output_text','text':text,'annotations':[]}]}
            response={'id':'resp_fixture_'+str(time.time_ns()),'object':'response','created_at':int(time.time()),'status':'completed','error':None,'incomplete_details':None,'model':payload.get('model','office-fixture'),'output':[item],'usage':{'input_tokens':50,'output_tokens':20,'total_tokens':70,'input_tokens_details':{'cached_tokens':0},'output_tokens_details':{'reasoning_tokens':0}}}
            print(json.dumps({'fixture':'responses','path':request_path,'model':response['model']}),flush=True)
            if not payload.get('stream'): return self.json(response)
            events=[('response.created',{'response':{**response,'status':'in_progress','output':[]}}),
                ('response.in_progress',{'response':{**response,'status':'in_progress','output':[]}}),
                ('response.output_item.added',{'output_index':0,'item':{**item,'status':'in_progress','content':[]}}),
                ('response.content_part.added',{'item_id':item['id'],'output_index':0,'content_index':0,'part':{'type':'output_text','text':'','annotations':[]}}),
                ('response.output_text.delta',{'item_id':item['id'],'output_index':0,'content_index':0,'delta':text}),
                ('response.output_text.done',{'item_id':item['id'],'output_index':0,'content_index':0,'text':text}),
                ('response.content_part.done',{'item_id':item['id'],'output_index':0,'content_index':0,'part':item['content'][0]}),
                ('response.output_item.done',{'output_index':0,'item':item}),('response.completed',{'response':response})]
            data=''.join('event: '+kind+'\ndata: '+json.dumps({'type':kind,'sequence_number':index,**value},ensure_ascii=False)+'\n\n' for index,(kind,value) in enumerate(events)).encode()
            self.send_response(200);self.send_header('Content-Type','text/event-stream');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data);self.wfile.flush()
            return
        if request_path in ('/v1/messages/count_tokens','/messages/count_tokens'):
            return self.json({'input_tokens':50})
        if request_path in ('/v1/messages','/messages'):
            print(json.dumps({'fixture':'anthropic','path':request_path,'model':payload.get('model','office-fixture')}),flush=True)
            # Native Claude protocol fixture; this is never a production model.
            text='隔离协议验收响应：已收到本轮请求。此响应来自模拟模型。'
            message={'id':'msg_fixture_'+str(time.time_ns()),'type':'message','role':'assistant',
                'model':payload.get('model','office-fixture'),'content':[{'type':'text','text':text}],
                'stop_reason':'end_turn','stop_sequence':None,'usage':{'input_tokens':50,'output_tokens':20}}
            if not payload.get('stream'):
                return self.json(message)
            events=[('message_start',{'type':'message_start','message':{**message,'content':[],'stop_reason':None,'usage':{'input_tokens':50,'output_tokens':0}}}),
                ('content_block_start',{'type':'content_block_start','index':0,'content_block':{'type':'text','text':''}}),
                ('content_block_delta',{'type':'content_block_delta','index':0,'delta':{'type':'text_delta','text':text}}),
                ('content_block_stop',{'type':'content_block_stop','index':0}),
                ('message_delta',{'type':'message_delta','delta':{'stop_reason':'end_turn','stop_sequence':None},'usage':{'output_tokens':20}}),
                ('message_stop',{'type':'message_stop'})]
            data=''.join('event: '+kind+'\ndata: '+json.dumps(value,ensure_ascii=False)+'\n\n' for kind,value in events).encode()
            self.send_response(200);self.send_header('Content-Type','text/event-stream');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data);self.wfile.flush()
            return
        messages=payload.get('messages',[])
        tool_names={tool.get('function',{}).get('name') for tool in payload.get('tools',[])}
        # Fixture diagnostics contain protocol shape only: no prompts, file
        # contents, request headers, URLs with credentials, or model secrets.
        print(json.dumps({'fixture':'office','path':self.path,'stream':bool(payload.get('stream')),
            'payload_keys':sorted(payload),'tool_names':sorted(str(name) for name in tool_names)},ensure_ascii=False),flush=True)
        if self.path not in ('/v1/chat/completions','/chat/completions'):
            return self.json({'error':{'message':'unsupported_fixture_protocol','type':'invalid_request_error'}},400)
        if not isinstance(payload.get('messages'),list):
            return self.json({'error':{'message':'chat_messages_required','type':'invalid_request_error'}},400)
        coding=any('CODING_ACCEPTANCE_FIXTURE' in str(message.get('content')) for message in messages)
        marker='coding-fixture-deliverables-created' if coding else 'office-fixture-deliverables-created'
        complete=any(message.get('role')=='tool' and marker in str(message.get('content')) for message in messages)
        call=None
        if 'bash' in tool_names and not complete:
            call={'index':0,'id':'fixture-create-office','type':'function','function':{'name':'bash','arguments':json.dumps({'command':CODE_SCRIPT if coding else SCRIPT,'description':'执行隔离验收任务','timeout':120000})}}
        text='已生成可编辑报告、分析表和演示稿。采购12、维护8，合计20；数据为隔离模拟验收资料。' if complete else '这是隔离模拟模型协议验收响应，未调用真实模型。'
        if complete and coding:text='已实际运行 Python 校验，12+8=20，生成代码补丁和验证说明。此任务使用隔离模拟模型。'
        identity={'id':'chatcmpl-office-fixture','object':'chat.completion.chunk','created':int(time.time()),'model':'office-fixture'}
        delta={'role':'assistant','tool_calls':[call]} if call else {'role':'assistant','content':text}
        finish='tool_calls' if call else 'stop'
        if payload.get('stream'):
            chunks=[{**identity,'choices':[{'index':0,'delta':delta,'finish_reason':None}]},
                    {**identity,'choices':[{'index':0,'delta':{},'finish_reason':finish}],'usage':{'prompt_tokens':50,'completion_tokens':20,'total_tokens':70}}]
            data=(''.join('data: '+json.dumps(chunk,ensure_ascii=False)+'\n\n' for chunk in chunks)+'data: [DONE]\n\n').encode()
            self.send_response(200);self.send_header('Content-Type','text/event-stream');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data);self.wfile.flush()
        else:
            message={'role':'assistant','content':None,'tool_calls':[dict((key,value) for key,value in call.items() if key!='index')]} if call else {'role':'assistant','content':text}
            self.json({**identity,'object':'chat.completion','choices':[{'index':0,'message':message,'finish_reason':finish}],'usage':{'prompt_tokens':50,'completion_tokens':20,'total_tokens':70}})


if __name__=='__main__':ThreadingHTTPServer(('0.0.0.0',8080),Handler).serve_forever()
