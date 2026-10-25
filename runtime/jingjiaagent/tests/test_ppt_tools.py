"""Regression checks for existing-PPT inspection and targeted copy editing."""
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.shapes.graphfrm import CT_GraphicalObjectFrame
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SKILLS_ROOT = Path(os.environ.get("JINGJIAAGENT_SKILLS_ROOT", str(ROOT / "skills")))
SPEC = importlib.util.spec_from_file_location(
    "inspect_deck", SKILLS_ROOT / "presentation/scripts/inspect_deck.py")
deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(deck)


class PresentationToolsTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "原始汇报.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(.2), Inches(5), Inches(.5)).text = "采购汇报"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
        for row, values in enumerate((("项目", "金额"), ("采购", "12"))):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        group = slide.shapes.add_group_shape()
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(.5)).text = "分组文本"
        data = CategoryChartData()
        data.categories = ["第一季度", "第二季度"]
        data.add_series("采购金额", (12, 18))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1),
                               Inches(4), Inches(3), data)
        second = presentation.slides.add_slide(presentation.slide_layouts[6])
        second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "保持不变"
        presentation.save(self.source)

    def test_inspection_locates_text_table_groups_and_chart(self):
        result = deck.inspect_deck(self.source)
        self.assertEqual(result["slide_count"], 2)
        self.assertGreater(result["slide_size_emu"]["width"], 0)
        shapes = result["slides"][0]["shapes"]
        self.assertEqual(shapes[0]["shape_path"], [0])
        self.assertEqual(shapes[0]["text"], "采购汇报")
        self.assertEqual(shapes[1]["table"]["cells"][1][1]["text"], "12")
        nested = shapes[2]["children"][0]["children"][0]
        self.assertEqual(nested["shape_path"], [2, 0, 0])
        self.assertEqual(nested["text"], "分组文本")
        self.assertEqual(shapes[3]["chart"]["series_names"], ["采购金额"])
        selected = deck.inspect_deck(self.source, 2)
        self.assertEqual([slide["slide_number"] for slide in selected["slides"]], [2])

    def test_modify_title_and_table_in_copy_preserves_source_and_other_content(self):
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        original = deck.inspect_deck(self.source)
        presentation = Presentation(self.source)
        shapes = presentation.slides[0].shapes
        # Use run-level edits to retain the existing formatting where possible.
        shapes[0].text_frame.paragraphs[0].runs[0].text = "采购汇报（修订）"
        shapes[1].table.cell(1, 1).text_frame.paragraphs[0].runs[0].text = "15"
        output = self.root / "汇报修订副本.pptx"
        presentation.save(output)
        changed = deck.inspect_deck(output)
        self.assertEqual(changed["slides"][0]["shapes"][0]["text"], "采购汇报（修订）")
        self.assertEqual(changed["slides"][0]["shapes"][1]["table"]["cells"][1][1]["text"], "15")
        self.assertEqual(changed["slides"][0]["shapes"][2:], original["slides"][0]["shapes"][2:])
        self.assertEqual(changed["slides"][1], original["slides"][1])
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)

    def test_smartart_keeps_location_and_other_slide_content(self):
        original = deck.inspect_deck(self.source)
        presentation = Presentation(self.source)
        slide = presentation.slides[0]
        frame = CT_GraphicalObjectFrame.new_graphicFrame(
            20, "SmartArt 流程图", Inches(1), Inches(4), Inches(4), Inches(2))
        frame.graphic.graphicData.set(
            "uri", "http://schemas.openxmlformats.org/drawingml/2006/diagram")
        slide.shapes._spTree.insert_element_before(frame, "p:extLst")
        presentation.save(self.source)
        before = self.source.read_bytes()

        result = deck.inspect_deck(self.source)
        shapes = result["slides"][0]["shapes"]
        self.assertEqual(shapes[:-1], original["slides"][0]["shapes"])
        self.assertEqual(result["slides"][1], original["slides"][1])
        unsupported = shapes[-1]
        self.assertEqual(unsupported["shape_path"], [4])
        self.assertEqual(unsupported["shape_id"], 20)
        self.assertEqual(unsupported["name"], "SmartArt 流程图")
        self.assertEqual(unsupported["bounds_emu"]["width"], Inches(4))
        self.assertEqual(unsupported["type"], "UNSUPPORTED")
        self.assertIn("未完整解析", unsupported["warnings"][0])
        self.assertEqual(self.source.read_bytes(), before)

    def test_unsupported_3d_chart_keeps_location_and_other_slide_content(self):
        original = deck.inspect_deck(self.source)
        presentation = Presentation(self.source)
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("采购", (12, 18))
        shape = presentation.slides[0].shapes.add_chart(
            XL_CHART_TYPE.PIE, Inches(1), Inches(4), Inches(4), Inches(2), data)
        shape.chart.has_title = True
        shape.chart.chart_title.text_frame.text = "采购占比"
        plot = shape.chart._chartSpace.xpath(".//c:pieChart")[0]
        # pie3DChart is a standard OOXML chart that python-pptx cannot inspect.
        plot.tag = qn("c:pie3DChart")
        for child in list(plot):
            if child.tag == qn("c:firstSliceAng"):
                plot.remove(child)
        presentation.save(self.source)
        before = self.source.read_bytes()

        result = deck.inspect_deck(self.source)
        shapes = result["slides"][0]["shapes"]
        self.assertEqual(shapes[:-1], original["slides"][0]["shapes"])
        self.assertEqual(result["slides"][1], original["slides"][1])
        unsupported = shapes[-1]
        self.assertEqual(unsupported["shape_path"], [4])
        self.assertEqual(unsupported["shape_id"], shape.shape_id)
        self.assertEqual(unsupported["name"], shape.name)
        self.assertEqual(unsupported["bounds_emu"]["width"], Inches(4))
        self.assertEqual(unsupported["chart"], {"type": "UNSUPPORTED", "title": "采购占比"})
        self.assertIn("未完整解析", unsupported["warnings"][0])
        self.assertEqual(self.source.read_bytes(), before)

    def test_malformed_chart_is_not_downgraded_to_warning(self):
        presentation = Presentation(self.source)
        chart = presentation.slides[0].shapes[3].chart
        chart._chartSpace.xpath(".//c:barDir")[0].set("val", "invalid-direction")
        presentation.save(self.source)
        with self.assertRaisesRegex(deck.InspectionError, "无法解析"):
            deck.inspect_deck(self.source)

    def test_cli_json_output_and_input_protection(self):
        output = self.root / "inspection" / "deck.json"
        self.assertEqual(deck.main([str(self.source), "--output", str(output)]), 0)
        self.assertEqual(json.loads(output.read_text(encoding="utf-8"))["slide_count"], 2)
        before = self.source.read_bytes()
        with contextlib.redirect_stderr(io.StringIO()) as stderr:
            status = deck.main([str(self.source), "--output", str(self.source)])
        self.assertEqual(status, 2)
        self.assertIn("不能覆盖", stderr.getvalue())
        self.assertEqual(self.source.read_bytes(), before)

    def test_bad_input_reports_clear_errors(self):
        for number in (0, 3):
            with self.subTest(number=number), self.assertRaisesRegex(deck.InspectionError, "页码超出"):
                deck.inspect_deck(self.source, number)
        broken = self.root / "broken.pptx"
        broken.write_bytes(b"not an office document")
        with contextlib.redirect_stderr(io.StringIO()) as stderr:
            self.assertEqual(deck.main([str(broken)]), 2)
        self.assertIn("损坏", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())
        with zipfile.ZipFile(broken, "w") as archive:
            archive.writestr("[Content_Types].xml", "<broken")
            archive.writestr("ppt/presentation.xml", "<broken")
        with self.assertRaisesRegex(deck.InspectionError, "无法解析"):
            deck.inspect_deck(broken)

    def test_input_and_output_budgets(self):
        for name, limit, message in (
            ("MAX_FILE_BYTES", 1, "大小限制"),
            ("MAX_EXPANDED_BYTES", 1, "解压后"),
            ("MAX_ARCHIVE_ENTRIES", 1, "过多文件"),
            ("MAX_TEXT_CHARS", 1, "文本内容过多"),
            ("MAX_TABLE_CELLS", 1, "单元格过多"),
            ("MAX_SHAPES", 1, "形状数量过多"),
        ):
            with self.subTest(limit=name), patch.object(deck, name, limit):
                with self.assertRaisesRegex(deck.InspectionError, message):
                    deck.inspect_deck(self.source)


if __name__ == "__main__":
    unittest.main()
