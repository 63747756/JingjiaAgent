"""Real offline rendering smoke. Run inside converter image, source mount read-only."""
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import zipfile

ROOT=Path('/src')


def load(name,path):
    spec=importlib.util.spec_from_file_location(name,path);value=importlib.util.module_from_spec(spec);spec.loader.exec_module(value);return value


def main():
    import openpyxl
    from docx import Document
    converter=load('converter',ROOT/'office/converter.py')
    # Prove the child filter prevents IPv4 and IPv6 even when the parent API
    # retains its internal listener. Test occurs in a disposable subprocess.
    source="import importlib.util,socket; s=importlib.util.spec_from_file_location('c','/src/office/converter.py');m=importlib.util.module_from_spec(s);s.loader.exec_module(m);m.deny_network_sockets();" + "\nfor family in (socket.AF_INET,socket.AF_INET6):\n try:socket.socket(family);raise RuntimeError('network permitted')\n except PermissionError:pass\n"
    subprocess.run([sys.executable,'-c',source],check=True)
    with tempfile.TemporaryDirectory(dir='/tmp') as directory:
        temporary=Path(directory)
        xlsx=temporary/'spreadsheet.xlsx';book=openpyxl.Workbook();sheet=book.active
        sheet.append(['项目','金额']);sheet.append(['采购',12]);sheet.append(['维护',8]);sheet.append(['合计','=SUM(B2:B3)']);book.save(xlsx)
        docx=temporary/'report.docx';doc=Document();doc.add_heading('中文工作报告',0);doc.add_paragraph('采购与维护费用合计20，数据来自用户提供的工作表。');doc.save(docx)
        fixtures=[('xlsx',xlsx),('docx',docx)]
        pptx=ROOT/'.build/office-smoke/presentation.pptx'
        if pptx.is_file():fixtures.append(('pptx',pptx))
        else:raise RuntimeError('generate PPTX skill fixture before smoke test')
        rendered_pdf=None
        for fmt,fixture in fixtures:
            request=temporary/fmt;request.mkdir();shutil.copyfile(fixture,request/('source.'+fmt))
            worker=subprocess.run([sys.executable,str(ROOT/'office/converter.py'),'--worker',str(request),fmt])
            if worker.returncode:
                print('fixture renderer diagnostic', (request/'renderer.stderr').read_text(errors='replace') if (request/'renderer.stderr').exists() else 'no stderr')
                raise RuntimeError((request/'error.json').read_text() if (request/'error.json').exists() else 'worker failed')
            with zipfile.ZipFile(request/'result.zip') as archive:
                validation=json.loads(archive.read('validation.json'))
                assert validation['status']=='passed',validation
                assert validation['page_count']>=1 and any(name.endswith('.png') for name in archive.namelist())
                assert archive.read('preview.pdf').startswith(b'%PDF-')
                if fmt=='xlsx':
                    calculated=openpyxl.load_workbook(request/'result/canonical.xlsx',data_only=True);assert calculated.active['B4'].value==20
                print(fmt,validation['status'],'pages='+str(validation['page_count']))
                if Path('/artifacts').is_dir():shutil.copytree(request/'result',Path('/artifacts')/fmt,dirs_exist_ok=True)
            if rendered_pdf is None:rendered_pdf=request/'result/preview.pdf'
        request=temporary/'pdf';request.mkdir();shutil.copyfile(rendered_pdf,request/'source.pdf')
        subprocess.run([sys.executable,str(ROOT/'office/converter.py'),'--worker',str(request),'pdf'],check=True)
        print('pdf passed; IPv4/IPv6 seccomp denied')
        analysis=load('analysis',ROOT/'skills/data-analysis/scripts/profile_data.py')
        csv=temporary/'data.csv';csv.write_text('名称,金额\n正常,12\n"=HYPERLINK(""http://bad"")",8\n',encoding='utf-8')
        result=analysis.profile(csv,temporary/'analysis');assert result['sheets'][0]['numeric']['金额']['sum']==20
        parsed=openpyxl.load_workbook(temporary/'analysis/analysis.xlsx');assert parsed.active['D3'].data_type=='s'
        report=load('report',ROOT/'skills/report-writing/scripts/build_report.py');report.build(json.loads((ROOT/'skills/report-writing/assets/report.template.json').read_text()),temporary/'report-generated.docx')
        check=load('check',ROOT/'skills/deliverable-check/scripts/check_deliverables.py');manifest=check.check(temporary/'analysis');assert len(manifest['files'])==2
        print('analysis/report/check deterministic scripts passed')
        # Exercise the actual HTTP entry point from the built image rather than
        # just the worker source mounted for diagnostics.
        token=temporary/'token';token.write_text('fixture-only-converter-token-0123456789abcdef')
        service=subprocess.Popen([sys.executable,'/opt/office/converter.py','--listen','127.0.0.1','--port','18090','--token-file',str(token)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        client=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        try:
            for attempt in range(30):
                try:client.open('http://127.0.0.1:18090/health',timeout=1).close();break
                except OSError:time.sleep(0.1)
            for authorization,expected in [('',401),('Bearer '+token.read_text(),200)]:
                request=urllib.request.Request('http://127.0.0.1:18090/convert?format=docx',data=docx.read_bytes(),headers={'Authorization':authorization,'Content-Type':'application/octet-stream'})
                try:
                    with client.open(request,timeout=70) as response:
                        assert response.status==expected
                        with zipfile.ZipFile(io.BytesIO(response.read())) as archive:
                            assert json.loads(archive.read('validation.json'))['status']=='passed'
                except urllib.error.HTTPError as error:assert error.code==expected
            print('built converter HTTP Bearer/ZIP protocol passed')
        finally:
            service.terminate();service.wait(timeout=5)


if __name__=='__main__':main()
