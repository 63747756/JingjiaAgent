"""Bounded, fixed-format Office rendering. No URL fetch or shell execution."""
import argparse
import hashlib
import hmac
import io
import json
import os
from pathlib import Path
import re
import signal
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit
import xml.etree.ElementTree as ET

MAX_INPUT = 32 << 20
MAX_EXPANDED = 128 << 20
MAX_OUTPUT = 128 << 20
MAX_PAGES = 200
TIMEOUT = 60
FORMATS = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}
SCHEMA = "jingjiaagent.office.validation.v1"
SLOTS = threading.BoundedSemaphore(1)


class ConversionError(ValueError):
    pass


def validate_input(source, fmt):
    """Recognize the requested document; reject archives that change scope."""
    if source.stat().st_size > MAX_INPUT:
        raise ConversionError("input_too_large")
    if fmt == "pdf":
        with source.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise ConversionError("invalid_pdf")
        return []
    if fmt not in FORMATS:
        raise ConversionError("unsupported_format")
    try:
        with zipfile.ZipFile(source) as archive:
            names = set()
            total = 0
            warnings = []
            if len(archive.infolist()) > 10000:
                raise ConversionError("too_many_package_entries")
            for entry in archive.infolist():
                name = entry.filename
                if (name in names or name.startswith("/") or "\\" in name or ":" in name
                        or any(part in ("..", ".", "") for part in name.rstrip("/").split("/"))
                        or any(value in name for value in ("\x00", "\r", "\n")) or ((entry.external_attr >> 16) & 0o170000) not in (0, 0o100000, 0o040000)):
                    raise ConversionError("unsafe_package_entry")
                names.add(name)
                total += entry.file_size
                if total > MAX_EXPANDED or entry.file_size > MAX_INPUT:
                    raise ConversionError("expanded_package_too_large")
                if "vbaproject" in name.lower() or name.lower().endswith(".bin") and "activex" in name.lower():
                    raise ConversionError("macros_or_activex_unsupported")
                if name.endswith((".xml", ".rels")):
                    markup = archive.read(entry)
                    # Reject local-file and network entity declarations before
                    # a native parser sees the package, including UTF-16 XML.
                    if re.search(rb"<!\s*(?:DOCTYPE|ENTITY)\b", markup.replace(b"\x00", b""), re.I):
                        raise ConversionError("external_xml_entities_unsupported")
                if name.endswith(".rels"):
                    if entry.file_size > (1 << 20):
                        raise ConversionError("relationship_file_too_large")
                    try:
                        relationships = ET.fromstring(markup)
                    except ET.ParseError as error:
                        raise ConversionError("invalid_relationship_xml") from error
                    for relationship in relationships.iter():
                        mode = relationship.attrib.get("TargetMode", "")
                        target = relationship.attrib.get("Target", "")
                        if mode.casefold() == "external" or urlsplit(target).scheme or "\\" in target:
                            # Even a file:// relationship could read the internal
                            # token file. Retain the original through the backend
                            # but never give these bytes to LibreOffice.
                            raise ConversionError("external_relationships_unsupported")
            if "[Content_Types].xml" not in names or FORMATS[fmt] not in names:
                raise ConversionError("format_content_mismatch")
            if archive.testzip() is not None:
                raise ConversionError("invalid_package_crc")
            return sorted(set(warnings))
    except zipfile.BadZipFile as error:
        raise ConversionError("invalid_office_package") from error


def fixed_process(arguments, deadline, cwd, stdout=None):
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise ConversionError("conversion_timeout")
    # The complete command is selected by the server. No submitted command,
    # filename, environment, URL, or shell syntax is ever executed.
    environment = {"PATH": "/usr/bin:/bin", "HOME": str(cwd), "TMPDIR": str(cwd),
                   "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8", "SAL_USE_VCLPLUGIN": "svp",
                   "SAL_DISABLE_OPENCL": "1", "MAX_CONCURRENCY": "1", "OMP_NUM_THREADS": "1",
                   "OPENBLAS_NUM_THREADS": "1"}
    try:
        with tempfile.TemporaryFile(dir=cwd) as output, (cwd / 'renderer.stderr').open('wb') as errors:
            result = subprocess.run(arguments, cwd=cwd, env=environment, timeout=remaining,
                                    stdin=subprocess.DEVNULL, stdout=stdout or output,
                                    stderr=errors, check=False)
            output.seek(0)
            captured = output.read(32768)
    except subprocess.TimeoutExpired as error:
        raise ConversionError("conversion_timeout") from error
    if result.returncode:
        raise ConversionError("renderer_failed")
    return captured


def convert(root, fmt):
    source = root / ("source." + fmt)
    warnings = validate_input(source, fmt)
    result = root / "result"
    result.mkdir()
    deadline = time.monotonic() + TIMEOUT - 3
    profile = root / "profile"
    profile.mkdir()
    (profile / "user").mkdir()
    # Disable document macros and automatic link updates in this fresh profile.
    (profile / "user/registrymodifications.xcu").write_text('''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>
</oor:items>''', encoding="utf-8")
    office = ["/usr/bin/libreoffice", "-env:UserInstallation=" + profile.as_uri(),
              "--headless", "--nologo", "--nodefault", "--norestore", "--nofirststartwizard"]
    input_for_pdf = source
    if fmt == "xlsx":
        canonical = root / "canonical"
        canonical.mkdir()
        fixed_process(office + ["--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(canonical), str(source)], deadline, root)
        input_for_pdf = canonical / "source.xlsx"
        if not input_for_pdf.is_file():
            raise ConversionError("recalculation_missing")
        import openpyxl
        formulas = openpyxl.load_workbook(input_for_pdf, read_only=True, data_only=False, keep_links=False)
        values = openpyxl.load_workbook(input_for_pdf, read_only=True, data_only=True, keep_links=False)
        checked = 0
        try:
            for sheet, value_sheet in zip(formulas.worksheets, values.worksheets):
                for formula_row, value_row in zip(sheet.iter_rows(), value_sheet.iter_rows()):
                    for formula, value in zip(formula_row, value_row):
                        checked += 1
                        if checked > 2_000_000:
                            raise ConversionError("spreadsheet_cell_limit")
                        if value.data_type == "e":
                            warnings.append("spreadsheet_formula_errors")
                        if formula.data_type == "f" and value.value is None:
                            warnings.append("spreadsheet_formula_without_cached_value")
        finally:
            formulas.close(); values.close()
        (result / "canonical.xlsx").write_bytes(input_for_pdf.read_bytes())
    if fmt == "pdf":
        (result / "preview.pdf").write_bytes(source.read_bytes())
    else:
        pdf_filter = {"docx": "pdf:writer_pdf_Export", "xlsx": "pdf:calc_pdf_Export", "pptx": "pdf:impress_pdf_Export"}[fmt]
        fixed_process(office + ["--convert-to", pdf_filter, "--outdir", str(result), str(input_for_pdf)], deadline, root)
        rendered = result / "source.pdf"
        if not rendered.is_file():
            raise ConversionError("preview_missing")
        rendered.rename(result / "preview.pdf")
    pdf = result / "preview.pdf"
    if pdf.stat().st_size > MAX_OUTPUT:
        raise ConversionError("preview_too_large")
    info = fixed_process(["/usr/bin/pdfinfo", str(pdf)], deadline, root).decode("utf-8", errors="replace")
    match = re.search(r"^Pages:\s+(\d+)\s*$", info, re.M)
    if not match or not 1 <= int(match.group(1)) <= MAX_PAGES:
        raise ConversionError("page_limit_or_invalid_pdf")
    page_count = int(match.group(1))
    pages = result / "pages"
    pages.mkdir()
    fixed_process(["/usr/bin/pdftoppm", "-png", "-r", "90", "-f", "1", "-l", str(page_count), str(pdf), str(pages / "page")], deadline, root, stdout=subprocess.DEVNULL)
    images = sorted(pages.glob("page-*.png"))
    if len(images) != page_count:
        raise ConversionError("page_render_incomplete")
    if sum(file.stat().st_size for file in result.rglob("*") if file.is_file()) > MAX_OUTPUT:
        raise ConversionError("output_too_large")
    warnings = sorted(set(warnings))
    validation = {"schema": SCHEMA, "status": "needs_review" if warnings else "passed", "format": fmt,
                  "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "page_count": page_count,
                  "warnings": warnings, "issues": warnings, "checks": {"package": True, "render": True, "page_images": True,
                  "spreadsheet_recalculated": fmt == "xlsx"}, "visual_review_required": True}
    (result / "validation.json").write_text(json.dumps(validation, ensure_ascii=False), encoding="utf-8")
    with zipfile.ZipFile(root / "result.zip", "w", compression=zipfile.ZIP_STORED) as archive:
        for file in sorted(result.rglob("*")):
            if file.is_file():
                archive.write(file, file.relative_to(result).as_posix())
    if (root / "result.zip").stat().st_size > MAX_OUTPUT:
        raise ConversionError("output_too_large")


def run_worker(root, fmt):
    import resource
    deny_network_sockets()
    resource.setrlimit(resource.RLIMIT_CPU, (TIMEOUT, TIMEOUT + 1))
    # LibreOffice reserves shared libraries/fonts and thread stacks in virtual
    # space. Actual resident memory is independently limited to 2 GiB by cgroup.
    resource.setrlimit(resource.RLIMIT_AS, (4 << 30, 4 << 30))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT, MAX_OUTPUT))
    resource.setrlimit(resource.RLIMIT_NOFILE, (128, 128))
    try:
        convert(root, fmt)
    except (ConversionError, OSError, ValueError) as error:
        (root / "error.json").write_text(json.dumps({"error": str(error)[:160]}), encoding="utf-8")
        return 1
    return 0


def deny_network_sockets():
    """Install a Linux seccomp filter inherited by LibreOffice and Poppler.

    Keep AF_UNIX for LibreOffice's local IPC; deny creation of every other
    socket family, including IPv4, IPv6 and packet sockets. Fail closed.
    """
    import ctypes
    import ctypes.util
    import errno
    class Compare(ctypes.Structure):
        _fields_ = [("arg", ctypes.c_uint), ("op", ctypes.c_int),
                    ("a", ctypes.c_uint64), ("b", ctypes.c_uint64)]
    library = ctypes.CDLL(ctypes.util.find_library("seccomp") or "libseccomp.so.2", use_errno=True)
    library.seccomp_init.argtypes = [ctypes.c_uint32]; library.seccomp_init.restype = ctypes.c_void_p
    library.seccomp_syscall_resolve_name.argtypes = [ctypes.c_char_p]; library.seccomp_syscall_resolve_name.restype = ctypes.c_int
    library.seccomp_rule_add_array.argtypes = [ctypes.c_void_p, ctypes.c_uint32, ctypes.c_int, ctypes.c_uint, ctypes.POINTER(Compare)]
    library.seccomp_load.argtypes = [ctypes.c_void_p]; library.seccomp_release.argtypes = [ctypes.c_void_p]
    context = library.seccomp_init(0x7fff0000)  # SCMP_ACT_ALLOW
    if not context: raise RuntimeError("network isolation unavailable")
    try:
        compare = Compare(0, 1, 1, 0)  # SCMP_CMP_NE, AF_UNIX
        syscall = library.seccomp_syscall_resolve_name(b"socket")
        if syscall < 0 or library.seccomp_rule_add_array(context, 0x00050000 | errno.EPERM, syscall, 1, ctypes.byref(compare)) or library.seccomp_load(context):
            raise RuntimeError("network isolation unavailable")
    finally:
        library.seccomp_release(context)


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    def log_message(self, *_args):
        pass  # No filenames, authorization tokens or document content in logs.

    def reply(self, code, value):
        data = json.dumps(value).encode()
        self.send_response(code); self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data))); self.send_header("Connection", "close")
        self.end_headers(); self.wfile.write(data); self.close_connection = True

    def do_GET(self):
        if self.path == "/health":
            return self.reply(200, {"ready": True, "schema": SCHEMA})
        return self.reply(404, {"error": "not_found"})

    def do_POST(self):
        token = self.server.token
        auth = self.headers.get("Authorization", "")
        if not token or not hmac.compare_digest(auth, "Bearer " + token) or self.headers.get("Origin"):
            return self.reply(401, {"error": "unauthorized"})
        parsed = urlsplit(self.path)
        query = parse_qs(parsed.query)
        fmt = query.get("format", [""])[0]
        if parsed.path != "/convert" or set(query) != {"format"} or len(query["format"]) != 1 or fmt not in (*FORMATS, "pdf"):
            return self.reply(400, {"error": "unsupported_format"})
        try:
            size = int(self.headers.get("Content-Length", ""))
        except ValueError:
            return self.reply(411, {"error": "content_length_required"})
        if self.headers.get("Transfer-Encoding") or not 0 < size <= MAX_INPUT:
            return self.reply(413, {"error": "input_limit"})
        if not SLOTS.acquire(blocking=False):
            return self.reply(503, {"error": "busy"})
        try:
            self.connection.settimeout(30)
            with tempfile.TemporaryDirectory(prefix="office-", dir="/tmp") as directory:
                root = Path(directory)
                with (root / ("source." + fmt)).open("wb") as stream:
                    remaining = size
                    while remaining:
                        chunk = self.rfile.read(min(65536, remaining))
                        if not chunk:
                            return self.reply(400, {"error": "incomplete_upload"})
                        stream.write(chunk); remaining -= len(chunk)
                process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), "--worker", str(root), fmt],
                    stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                    start_new_session=True)
                try:
                    status = process.wait(timeout=TIMEOUT)
                except subprocess.TimeoutExpired:
                    os.killpg(process.pid, signal.SIGKILL); process.wait()
                    return self.reply(422, {"error": "conversion_timeout"})
                finally:
                    # LibreOffice occasionally keeps a child after its launcher
                    # exits; the request process group is always reclaimed.
                    try: os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError: pass
                archive = root / "result.zip"
                if status or not archive.is_file():
                    error = json.loads((root / "error.json").read_text()) if (root / "error.json").is_file() else {"error": "renderer_failed"}
                    return self.reply(422, error)
                self.send_response(200); self.send_header("Content-Type", "application/zip")
                self.send_header("Content-Length", str(archive.stat().st_size)); self.send_header("Cache-Control", "no-store")
                self.end_headers()
                with archive.open("rb") as stream:
                    while chunk := stream.read(65536): self.wfile.write(chunk)
        except (OSError, TimeoutError, ValueError):
            self.close_connection = True
        finally:
            SLOTS.release()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", nargs=2)
    parser.add_argument("--listen", default="0.0.0.0")
    parser.add_argument("--port", type=int, default=8090)
    parser.add_argument("--token-file", default="/run/secrets/converter_token")
    args = parser.parse_args()
    if args.worker:
        return run_worker(Path(args.worker[0]), args.worker[1])
    token = Path(args.token_file).read_text().strip()
    if len(token) < 32:
        raise SystemExit("converter token must contain at least 32 characters")
    if os.geteuid() == 0:
        # Compose file-backed secrets retain the installer's UID/0600 mode.
        # Read the secret before opening a listener, then permanently remove
        # root identity and all effective capabilities. Workers inherit 10001.
        os.setgroups([])
        os.setgid(10001)
        os.setuid(10001)
        if os.getuid() != 10001 or os.geteuid() != 10001:
            raise SystemExit("converter privilege drop failed")
    server = ThreadingHTTPServer((args.listen, args.port), Handler)
    server.daemon_threads = True
    server.token = token
    server.serve_forever()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
